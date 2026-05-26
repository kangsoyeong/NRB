import openpyxl

FILE1 = "/root/.claude/uploads/24041d56-69af-4927-919e-4cbd89e59ed0/35a33838-_____260523.xlsx"
FILE2 = "/root/.claude/uploads/24041d56-69af-4927-919e-4cbd89e59ed0/875360de-__________.xlsx"
OUTPUT = "/root/.claude/uploads/24041d56-69af-4927-919e-4cbd89e59ed0/result_matched.xlsx"

# ── 1. Inspect headers ──────────────────────────────────────────────────────

print("=" * 60)
print("FILE 1  –  'Cash out' sheet  (row 2 = header)")
wb1 = openpyxl.load_workbook(FILE1, data_only=True)
ws1 = wb1["Cash out"]

header_row1 = list(ws1.iter_rows(min_row=2, max_row=2, values_only=True))[0]
for i, h in enumerate(header_row1):
    if h is not None:
        print(f"  col {i:>3d}  ({chr(65+i) if i < 26 else '?'})  {h}")

print()
print("FILE 2  –  'Sheet1'  (row 1 = header)")
wb2 = openpyxl.load_workbook(FILE2)
ws2 = wb2["Sheet1"]

header_row2 = list(ws2.iter_rows(min_row=1, max_row=1, values_only=True))[0]
for i, h in enumerate(header_row2):
    if h is not None:
        print(f"  col {i:>3d}  ({chr(65+i) if i < 26 else '?'})  {h}")

# ── 2. Locate the exact column indices ─────────────────────────────────────

def find_col(headers, *names):
    """Return 0-based index of the first header that matches any of 'names'."""
    for i, h in enumerate(headers):
        if h is not None and str(h).strip() in names:
            return i
    return None

# File 1 – expected B(1) = 지결문서번호, Q(16) = 원천전표번호
COL1_DOC   = find_col(header_row1, "지결문서번호")
COL1_SRC   = find_col(header_row1, "원천전표번호")

if COL1_DOC is None:
    COL1_DOC = 1   # fallback per task spec
    print(f"\n[WARN] '지결문서번호' not found in File1 headers – using fallback col {COL1_DOC}")
if COL1_SRC is None:
    COL1_SRC = 16
    print(f"[WARN] '원천전표번호' not found in File1 headers – using fallback col {COL1_SRC}")

# File 2 – expected O(14) = 지결문서번호, P(15) = 원천전표번호
COL2_DOC   = find_col(header_row2, "지결문서번호")
COL2_SRC   = find_col(header_row2, "원천전표번호")

if COL2_DOC is None:
    COL2_DOC = 14
    print(f"[WARN] '지결문서번호' not found in File2 headers – using fallback col {COL2_DOC}")
if COL2_SRC is None:
    COL2_SRC = 15
    print(f"[WARN] '원천전표번호' not found in File2 headers – using fallback col {COL2_SRC}")

print()
print(f"File1 mapping columns  →  doc={COL1_DOC} ({chr(65+COL1_DOC)}),  src={COL1_SRC} ({chr(65+COL1_SRC)})")
print(f"File2 target columns   →  doc={COL2_DOC} ({chr(65+COL2_DOC)}),  src={COL2_SRC} ({chr(65+COL2_SRC)})")

# ── 3. Build lookup mapping from File 1 ────────────────────────────────────

mapping = {}   # {지결문서번호: [원천전표번호, ...]}

for row in ws1.iter_rows(min_row=3, values_only=True):   # data from row 3
    doc = row[COL1_DOC] if len(row) > COL1_DOC else None
    src = row[COL1_SRC] if len(row) > COL1_SRC else None

    if doc is None or str(doc).strip() == "":
        continue
    if src is None or str(src).strip() == "":
        continue

    doc_key = str(doc).strip()
    src_val = str(src).strip()

    if doc_key not in mapping:
        mapping[doc_key] = []
    if src_val not in mapping[doc_key]:
        mapping[doc_key].append(src_val)

print(f"\nLookup map built: {len(mapping)} unique 지결문서번호 entries")
# Show a few samples
for k, v in list(mapping.items())[:5]:
    print(f"  {k!r:40s}  →  {v}")

# ── 4. Fill blanks in File 2 ───────────────────────────────────────────────

total_rows   = 0
blank_before = 0
filled       = 0
still_blank  = 0

# openpyxl rows are 1-based; header is row 1, data from row 2
for row_cells in ws2.iter_rows(min_row=2):
    total_rows += 1

    doc_cell = row_cells[COL2_DOC]
    src_cell = row_cells[COL2_SRC]

    doc_val = doc_cell.value
    src_val = src_cell.value

    # Determine if src is blank
    src_blank = (src_val is None or str(src_val).strip() == "")

    if src_blank:
        blank_before += 1
        if doc_val is not None and str(doc_val).strip() != "":
            doc_key = str(doc_val).strip()
            if doc_key in mapping:
                src_cell.value = mapping[doc_key][0]   # use first match
                filled += 1
            else:
                still_blank += 1
        else:
            still_blank += 1

# ── 5. Save ────────────────────────────────────────────────────────────────

wb2.save(OUTPUT)
print(f"\nSaved  →  {OUTPUT}")

# ── 6. Summary ─────────────────────────────────────────────────────────────

print()
print("=" * 60)
print("SUMMARY")
print("=" * 60)
print(f"  Total rows in File 2            : {total_rows}")
print(f"  Rows with blank 원천전표번호 (before) : {blank_before}")
print(f"  Successfully filled             : {filled}")
print(f"  Still blank after matching      : {still_blank}")
print(f"  Already had a value             : {total_rows - blank_before}")
