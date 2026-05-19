"""
NRB 차입금현황 기반 잔액 및 은행별 가중평균금리 계산
입력: 차입금현황 xlsx (1-1. 차입금현황 시트)
출력: 잔액_가중평균금리.xlsx
"""

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import datetime
import sys
import os

SOURCE_FILE = "/root/.claude/uploads/cd7b7295-6998-41b8-9d78-afbea185df46/fce45296-_______26.05.19.xlsx"
OUTPUT_FILE = os.path.join(os.path.dirname(__file__), "잔액_가중평균금리.xlsx")

BANKS = ["기업은행", "광주은행", "전북은행", "신협은행", "우리은행", "중소벤처기업진흥공단"]
SKIP_TYPES = {"이자", "수입신용장"}
DATA_ROW_START = 18
DATA_ROW_END = 88
HEADER_ROW = 17
MONTH_COL_START = 15  # O열 (2025-12-31)


def load_loans(ws):
    """원금 차입금 행만 추출 (이자·수입신용장 제외)."""
    month_cols = {}
    for col in range(MONTH_COL_START, ws.max_column + 1):
        v = ws.cell(row=HEADER_ROW, column=col).value
        if isinstance(v, datetime.datetime):
            month_cols[col] = v

    loans = []
    for row in range(DATA_ROW_START, DATA_ROW_END + 1):
        구분 = ws.cell(row=row, column=2).value
        은행명 = ws.cell(row=row, column=3).value
        if not 구분 or 구분 in SKIP_TYPES or 은행명 not in BANKS:
            continue
        rate = ws.cell(row=row, column=13).value
        if not isinstance(rate, (int, float)):
            continue
        balances = {}
        for col, dt in month_cols.items():
            val = ws.cell(row=row, column=col).value
            if isinstance(val, (int, float)):
                balances[dt] = val
        loans.append({"구분": 구분, "은행명": 은행명, "이자율": rate, "balances": balances})

    return loans, sorted(month_cols.values())


def compute_bank_stats(loans, months):
    """월별 은행별 잔액 및 가중평균금리 계산."""
    result = {}  # {month: {bank: {"balance": x, "rate": y}}}
    for month in months:
        result[month] = {}
        for bank in BANKS:
            total_balance = 0
            weighted_sum = 0.0
            for loan in loans:
                if loan["은행명"] == bank and month in loan["balances"]:
                    bal = loan["balances"][month]
                    total_balance += bal
                    weighted_sum += bal * loan["이자율"]
            if total_balance > 0:
                result[month][bank] = {
                    "balance": total_balance,
                    "rate": weighted_sum / total_balance,
                }
    return result


def compute_total_stats(loans, months):
    """월별 전체 합계 및 가중평균금리 계산."""
    result = {}
    for month in months:
        total_balance = 0
        weighted_sum = 0.0
        for bank in BANKS:
            for loan in loans:
                if loan["은행명"] == bank and month in loan["balances"]:
                    bal = loan["balances"][month]
                    total_balance += bal
                    weighted_sum += bal * loan["이자율"]
        if total_balance > 0:
            result[month] = {"balance": total_balance, "rate": weighted_sum / total_balance}
    return result


# ─── 스타일 ────────────────────────────────────────────────────────────────────

thin = Side(style="thin", color="000000")
medium = Side(style="medium", color="000000")


def border_all(thick=False):
    s = medium if thick else thin
    return Border(left=s, right=s, top=s, bottom=s)


HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F3864")
SUBHEADER_FILL = PatternFill(fill_type="solid", fgColor="2E74B5")
TOTAL_FILL = PatternFill(fill_type="solid", fgColor="D6E4F0")
ALT_FILL = PatternFill(fill_type="solid", fgColor="F2F7FC")

WHITE_BOLD = Font(name="맑은 고딕", bold=True, color="FFFFFF", size=10)
WHITE_NORM = Font(name="맑은 고딕", color="FFFFFF", size=10)
BLACK_NORM = Font(name="맑은 고딕", size=10)
BLACK_BOLD = Font(name="맑은 고딕", bold=True, size=10)
TOTAL_FONT = Font(name="맑은 고딕", bold=True, size=10, color="1F3864")

CENTER = Alignment(horizontal="center", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")


def apply(cell, value=None, font=None, fill=None, border=None, alignment=None, number_format=None):
    if value is not None:
        cell.value = value
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if border:
        cell.border = border
    if alignment:
        cell.alignment = alignment
    if number_format:
        cell.number_format = number_format


def write_excel(stats, total_stats, months, output_path):
    wb = openpyxl.Workbook()

    # ── 시트 1: 잔액 ────────────────────────────────────────────────────────────
    ws_bal = wb.active
    ws_bal.title = "잔액"

    ws_bal.row_dimensions[1].height = 30
    ws_bal.row_dimensions[2].height = 22

    # 제목
    ws_bal.merge_cells("A1:A2")
    apply(ws_bal["A1"], "구분", WHITE_BOLD, HEADER_FILL, border_all(), CENTER)

    col_offset = 2
    for i, month in enumerate(months):
        c = get_column_letter(col_offset + i)
        label = month.strftime("%Y-%m")
        apply(ws_bal[f"{c}1"], label, WHITE_BOLD, HEADER_FILL, border_all(), CENTER)
        ws_bal.column_dimensions[c].width = 18

    ws_bal.column_dimensions["A"].width = 20

    # 은행별 데이터
    bank_rows = {}
    for b_idx, bank in enumerate(BANKS):
        r = 3 + b_idx
        bank_rows[bank] = r
        ws_bal.row_dimensions[r].height = 18
        fill = ALT_FILL if b_idx % 2 else PatternFill()
        apply(ws_bal.cell(row=r, column=1), bank, BLACK_BOLD, fill, border_all(), LEFT)
        for i, month in enumerate(months):
            cell = ws_bal.cell(row=r, column=col_offset + i)
            val = stats.get(month, {}).get(bank, {}).get("balance")
            apply(cell, val, BLACK_NORM, fill, border_all(), RIGHT, '#,##0')

    # 합계 행
    total_row = 3 + len(BANKS)
    ws_bal.row_dimensions[total_row].height = 20
    apply(ws_bal.cell(row=total_row, column=1), "합계", TOTAL_FONT, TOTAL_FILL, border_all(), LEFT)
    for i, month in enumerate(months):
        cell = ws_bal.cell(row=total_row, column=col_offset + i)
        val = total_stats.get(month, {}).get("balance")
        apply(cell, val, TOTAL_FONT, TOTAL_FILL, border_all(), RIGHT, '#,##0')

    # ── 시트 2: 가중평균금리 ────────────────────────────────────────────────────
    ws_rate = wb.create_sheet("은행별 가중평균금리")

    ws_rate.row_dimensions[1].height = 30
    ws_rate.row_dimensions[2].height = 22
    ws_rate.column_dimensions["A"].width = 20

    apply(ws_rate["A1"], "구분", WHITE_BOLD, HEADER_FILL, border_all(), CENTER)
    ws_rate.merge_cells("A1:A2")

    for i, month in enumerate(months):
        c = get_column_letter(col_offset + i)
        label = month.strftime("%Y-%m")
        apply(ws_rate[f"{c}1"], label, WHITE_BOLD, HEADER_FILL, border_all(), CENTER)
        ws_rate.column_dimensions[c].width = 14

    for b_idx, bank in enumerate(BANKS):
        r = 3 + b_idx
        ws_rate.row_dimensions[r].height = 18
        fill = ALT_FILL if b_idx % 2 else PatternFill()
        apply(ws_rate.cell(row=r, column=1), bank, BLACK_BOLD, fill, border_all(), LEFT)
        for i, month in enumerate(months):
            cell = ws_rate.cell(row=r, column=col_offset + i)
            info = stats.get(month, {}).get(bank)
            val = info["rate"] if info else None
            apply(cell, val, BLACK_NORM, fill, border_all(), CENTER, '0.0000%')

    # 전체 가중평균금리 합계
    total_row = 3 + len(BANKS)
    ws_rate.row_dimensions[total_row].height = 20
    apply(ws_rate.cell(row=total_row, column=1), "전체 가중평균금리", TOTAL_FONT, TOTAL_FILL, border_all(), LEFT)
    for i, month in enumerate(months):
        cell = ws_rate.cell(row=total_row, column=col_offset + i)
        info = total_stats.get(month)
        val = info["rate"] if info else None
        apply(cell, val, TOTAL_FONT, TOTAL_FILL, border_all(), CENTER, '0.0000%')

    # ── 시트 3: 통합 (잔액 + 금리) ──────────────────────────────────────────────
    ws_comb = wb.create_sheet("통합")

    ws_comb.row_dimensions[1].height = 30
    ws_comb.column_dimensions["A"].width = 20
    ws_comb.column_dimensions["B"].width = 14

    apply(ws_comb["A1"], "은행", WHITE_BOLD, HEADER_FILL, border_all(), CENTER)
    apply(ws_comb["B1"], "구분", WHITE_BOLD, HEADER_FILL, border_all(), CENTER)

    for i, month in enumerate(months):
        c = get_column_letter(3 + i)
        label = month.strftime("%Y-%m")
        apply(ws_comb[f"{c}1"], label, WHITE_BOLD, HEADER_FILL, border_all(), CENTER)
        ws_comb.column_dimensions[c].width = 16

    r = 2
    for b_idx, bank in enumerate(BANKS):
        fill = ALT_FILL if b_idx % 2 else PatternFill()

        # 잔액 행
        ws_comb.row_dimensions[r].height = 18
        apply(ws_comb.cell(row=r, column=1), bank, BLACK_BOLD, fill, border_all(), LEFT)
        apply(ws_comb.cell(row=r, column=2), "잔액(원)", BLACK_NORM, fill, border_all(), CENTER)
        for i, month in enumerate(months):
            cell = ws_comb.cell(row=r, column=3 + i)
            val = stats.get(month, {}).get(bank, {}).get("balance")
            apply(cell, val, BLACK_NORM, fill, border_all(), RIGHT, '#,##0')

        # 금리 행
        r += 1
        ws_comb.row_dimensions[r].height = 18
        apply(ws_comb.cell(row=r, column=1), "", BLACK_NORM, fill, border_all(), CENTER)
        apply(ws_comb.cell(row=r, column=2), "가중평균금리", BLACK_NORM, fill, border_all(), CENTER)
        for i, month in enumerate(months):
            cell = ws_comb.cell(row=r, column=3 + i)
            info = stats.get(month, {}).get(bank)
            val = info["rate"] if info else None
            apply(cell, val, BLACK_NORM, fill, border_all(), CENTER, '0.0000%')
        r += 1

    # 합계 행
    ws_comb.row_dimensions[r].height = 20
    apply(ws_comb.cell(row=r, column=1), "합계", TOTAL_FONT, TOTAL_FILL, border_all(), LEFT)
    apply(ws_comb.cell(row=r, column=2), "잔액(원)", TOTAL_FONT, TOTAL_FILL, border_all(), CENTER)
    for i, month in enumerate(months):
        cell = ws_comb.cell(row=r, column=3 + i)
        val = total_stats.get(month, {}).get("balance")
        apply(cell, val, TOTAL_FONT, TOTAL_FILL, border_all(), RIGHT, '#,##0')

    r += 1
    ws_comb.row_dimensions[r].height = 20
    apply(ws_comb.cell(row=r, column=1), "", TOTAL_FONT, TOTAL_FILL, border_all(), CENTER)
    apply(ws_comb.cell(row=r, column=2), "전체 가중평균금리", TOTAL_FONT, TOTAL_FILL, border_all(), CENTER)
    for i, month in enumerate(months):
        cell = ws_comb.cell(row=r, column=3 + i)
        info = total_stats.get(month)
        val = info["rate"] if info else None
        apply(cell, val, TOTAL_FONT, TOTAL_FILL, border_all(), CENTER, '0.0000%')

    wb.save(output_path)
    print(f"저장 완료: {output_path}")


def main():
    src = SOURCE_FILE
    if not os.path.exists(src):
        print(f"[오류] 파일 없음: {src}", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb["1-1. 차입금현황"]

    loans, months = load_loans(ws)

    # 2025-12-31 제외 (기준 잔액이므로 분석 대상에서 제외)
    months = [m for m in months if m.year >= 2026]

    stats = compute_bank_stats(loans, months)
    total_stats = compute_total_stats(loans, months)

    write_excel(stats, total_stats, months, OUTPUT_FILE)

    # 터미널 요약 출력
    print("\n=== 잔액 및 은행별 가중평균금리 요약 (2026년) ===")
    months_2026 = [m for m in months if m.year == 2026]
    header = f"{'은행':<16}" + "".join(f"{m.strftime('%Y-%m'):>12}" for m in months_2026)
    print(header)
    print("-" * len(header))
    for bank in BANKS:
        line = f"{bank:<16}"
        for m in months_2026:
            info = stats.get(m, {}).get(bank)
            if info:
                line += f"{info['rate']:>11.4%} "
            else:
                line += f"{'':>12}"
        print(line)
    print("-" * len(header))
    line = f"{'전체 가중평균':<16}"
    for m in months_2026:
        info = total_stats.get(m)
        if info:
            line += f"{info['rate']:>11.4%} "
        else:
            line += f"{'':>12}"
    print(line)

    print("\n=== 잔액 요약 (2026년, 단위: 억원) ===")
    print(header)
    print("-" * len(header))
    for bank in BANKS:
        line = f"{bank:<16}"
        for m in months_2026:
            info = stats.get(m, {}).get(bank)
            if info:
                line += f"{info['balance']/1e8:>11,.1f} "
            else:
                line += f"{'':>12}"
        print(line)
    print("-" * len(header))
    line = f"{'합계':<16}"
    for m in months_2026:
        info = total_stats.get(m)
        if info:
            line += f"{info['balance']/1e8:>11,.1f} "
        else:
            line += f"{'':>12}"
    print(line)


if __name__ == "__main__":
    main()
