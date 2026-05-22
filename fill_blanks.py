import openpyxl
from openpyxl.styles import numbers
import copy
import shutil

src = '/root/.claude/uploads/6f4d879a-ce9c-4c4c-91e7-97f86c4f246d/7750a3ec-_____________.xlsx'
dst = '/home/user/NRB/의왕초평_도급기성청구현황.xlsx'

shutil.copy2(src, dst)

wb = openpyxl.load_workbook(dst)
ws = wb.active

# 회차행, 공장행, 현장행 index (1-based)
groups = [
    (3, 4, 5),
    (6, 7, 8),
    (9, 10, 11),
    (12, 13, 14),
    (15, 16, 17),
]

COMP_COLS = [4, 5, 6, 7, 8]  # D~H: 극동건설, HJ, 서한, 대저, 엔알비

def copy_style(src_cell, dst_cell):
    dst_cell.font          = copy.copy(src_cell.font)
    dst_cell.fill          = copy.copy(src_cell.fill)
    dst_cell.border        = copy.copy(src_cell.border)
    dst_cell.alignment     = copy.copy(src_cell.alignment)
    dst_cell.number_format = '#,##0'

for (r_main, r_fac, r_fld) in groups:
    fac   = ws.cell(row=r_fac,  column=2).value
    fld   = ws.cell(row=r_fld,  column=2).value

    if not fac or not fld:
        continue
    total = fac + fld   # 회차 합계셀이 수식이므로 직접 계산

    for col in COMP_COLS:
        comp_total = ws.cell(row=r_main, column=col).value
        if not comp_total:
            continue

        f  = round(comp_total * fac / total)
        fi = comp_total - f

        c_fac = ws.cell(row=r_fac, column=col, value=f)
        copy_style(ws.cell(row=r_main, column=col), c_fac)

        c_fld = ws.cell(row=r_fld, column=col, value=fi)
        copy_style(ws.cell(row=r_main, column=col), c_fld)

wb.save(dst)
print("저장 완료:", dst)

# 검증
wb2 = openpyxl.load_workbook(dst, data_only=True)
ws2 = wb2.active
print()
for (r_main, r_fac, r_fld) in groups:
    label    = ws2.cell(row=r_main, column=1).value
    fac_tgt  = ws2.cell(row=r_fac,  column=2).value
    fld_tgt  = ws2.cell(row=r_fld,  column=2).value
    fac_sum  = sum((ws2.cell(row=r_fac, column=c).value or 0) for c in COMP_COLS)
    fld_sum  = sum((ws2.cell(row=r_fld, column=c).value or 0) for c in COMP_COLS)
    ok = "✓" if fac_sum == fac_tgt and fld_sum == fld_tgt else "✗"
    print(f"{ok} {label}: 공장 업체합={fac_sum:,} / 목표={fac_tgt:,}  |  현장 업체합={fld_sum:,} / 목표={fld_tgt:,}")
