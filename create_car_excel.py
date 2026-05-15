import openpyxl
from openpyxl.styles import (
    Font, Alignment, PatternFill, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "차량"

# 스타일 정의
thin = Side(style='thin', color='000000')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

header_fill = PatternFill(fill_type='solid', fgColor='1F3864')   # 진한 남색
subheader_fill = PatternFill(fill_type='solid', fgColor='2E74B5') # 중간 파란색
total_fill = PatternFill(fill_type='solid', fgColor='D6E4F0')     # 연한 파란색
changed_fill = PatternFill(fill_type='solid', fgColor='FFF2CC')   # 노란색 (변경)

white_font = Font(name='맑은 고딕', bold=True, color='FFFFFF', size=10)
black_font = Font(name='맑은 고딕', size=10)
bold_black = Font(name='맑은 고딕', bold=True, size=10)
total_font = Font(name='맑은 고딕', bold=True, size=10, color='1F3864')

center = Alignment(horizontal='center', vertical='center', wrap_text=True)
right_align = Alignment(horizontal='right', vertical='center')

# 열 너비 설정
col_widths = [12, 12, 16, 14, 13, 12, 12, 8, 13, 12, 14, 13, 12]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# 행 높이
ws.row_dimensions[1].height = 18
ws.row_dimensions[2].height = 28
ws.row_dimensions[3].height = 36

# ── 1행: 제목 ──────────────────────────────────────────
ws.merge_cells('A1:M1')
c = ws['A1']
c.value = '■ 차량'
c.font = Font(name='맑은 고딕', bold=True, size=11)
c.alignment = Alignment(horizontal='left', vertical='center')

# ── 2행: 대분류 헤더 ──────────────────────────────────
headers_row2 = {
    'A2': '구분', 'B2': '하이패스', 'C2': '부서', 'D2': '차량명',
    'E2': '차량번호',
}
ws.merge_cells('F2:H2')
ws['F2'].value = '계약기간'
ws.merge_cells('I2:I3')
ws['I2'].value = '보증금'
ws.merge_cells('J2:L2')
ws['J2'].value = '5월 청구 금액'
ws.merge_cells('M2:M3')
ws['M2'].value = '렌탈사'

for coord, val in headers_row2.items():
    r, c_num = int(coord[1:]), openpyxl.utils.column_index_from_string(coord[0])
    ws.merge_cells(f'{coord[0]}2:{coord[0]}3')
    cell = ws[coord]
    cell.value = val
    cell.font = white_font
    cell.fill = header_fill
    cell.alignment = center
    cell.border = border

for coord in ['F2', 'I2', 'J2', 'M2']:
    cell = ws[coord]
    cell.font = white_font
    cell.fill = header_fill
    cell.alignment = center
    cell.border = border

# ── 3행: 소분류 헤더 ──────────────────────────────────
sub_headers = {
    'F3': '시작', 'G3': '종료', 'H3': '지급일',
    'J3': '월 렌트료', 'K3': '월 기타비용\n(범칙금, 하이패스\n스료 미납)', 'L3': '총계'
}
for coord, val in sub_headers.items():
    cell = ws[coord]
    cell.value = val
    cell.font = white_font
    cell.fill = subheader_fill
    cell.alignment = center
    cell.border = border

# 2행 나머지 셀 스타일
for col in ['A', 'B', 'C', 'D', 'E']:
    ws[f'{col}3'].border = border

for col in ['I', 'M']:
    ws[f'{col}3'].border = border

# ── 데이터 ────────────────────────────────────────────
# 컬럼: 구분, 하이패스, 부서, 차량명, 차량번호, 시작, 종료, 지급일, 보증금, 월렌트료, 기타비용, 총계, 렌탈사
data = [
    ('강건우',      'O',         '대표이사',       '제네시스 G90',   '220하9211', '2025-12-05','2029-12-02','21일', 30000000, 1666060,  32000,  1698060, 'KB캐피탈'),
    ('서울공용1',   '주현철',    '경영기획실',     '아반떼',         '222허3499', '2024-02-13','2028-02-13','25일', None,      457000,   None,   457000,  '롯데렌터카'),
    ('배용덕',      '배용덕',    '영업관리실',     '그랜저 GN7',    '179호4893', '2025-03-21','2029-03-21','25일', None,      789000,   None,   789000,  '롯데렌터카'),
    ('군산공용1',   '박진우',    '모듈러생산실',   '스포티지 HEV',  '198호2160', '2023-02-04','2027-02-04','31일', None,      762740,   None,   762740,  '현대캐피탈'),
    ('군산공용2',   '군산직원전원','모듈러생산실', '포터',           '813수4246', '2024-12-23','2027-12-22','25일', None,      463027,   86973,  550000,  '롯데오토리스'),
    ('김성기',      '김성기',    '모듈러생산실',   '투싼 NX4',      '124호7101', '2022-09-24','2026-09-24','20일', None,      629310,   None,   629310,  '현대캐피탈'),
    ('최대현',      '최대현',    'PC생산실',       '그랜저 2.5',    '198호5335', '2023-02-27','2027-02-27','25일', None,      902110,   13700,  915810,  '현대캐피탈'),
    ('최진달미',    '최진달미',  'R&D본부',        '그랜저 2.5',    '200호3819', '2023-12-26','2027-12-26','25일', None,      843700,   None,   843700,  '현대캐피탈'),
    # 200호3807: 경영기획실 공용차량으로 변경 (노란색 표시)
    ('서울공용2',   'O',         '경영기획실',     '그랜저 2.5',    '200호3807', '2023-12-26','2027-12-26','25일', None,      843700,   None,   843700,  '현대캐피탈'),
    ('김준식',      '김준식',    '모듈러생산실',   '그랜저 2.5',    '200호2723', '2023-12-27','2027-12-27','25일', None,      843700,   5300,   849000,  '현대캐피탈'),
    # 200호2725 (오진택) 계약해지 -> 제외
    ('군산공용3',   '김효석',    '모듈러생산실',   '카니발',         '232수1531', '2022-03-30','2027-03-29','20일', None,      539923,   None,   539923,  'NH캐피탈'),
    ('김갑득',      '김갑득',    'R&D본부',        '제네시스 G80',  '222하6305', '2024-04-17','2028-04-17','25일', 5000000,   1140000,  None,   1140000, '롯데렌터카'),
    ('김도현',      '김도현',    '모듈러생산실',   '카니발 KA4',    '222호8286', '2024-09-13','2028-09-13','25일', None,      656000,   None,   656000,  '롯데렌터카'),
    ('김종록',      '김종록',    '모듈러생산실',   '투싼',           '230한1376', '2025-01-14','2029-01-14','25일', None,      640000,   None,   640000,  '롯데렌터카'),
    ('정윤영',      '정윤영',    '건축기술지원실', '그랜저 GN7',    '179호4610', '2025-02-24','2029-02-24','25일', None,      789000,   None,   789000,  '롯데렌터카'),
    ('김명진',      'O',         '경영기획실',     'K8 PE(2.5)',     '200허5988', '2025-06-02','2029-06-02','25일', None,      754600,   None,   754600,  '롯데렌터카'),
    ('이종명',      '이종명',    '모듈러생산실',   '그랜저 GN7',    '231호1562', '2025-07-02','2029-07-02','25일', None,      730000,   None,   730000,  '롯데렌터카'),
]

# 변경된 행 인덱스 (0-based, 데이터 리스트 기준)
changed_row_idx = 8  # 200호3807 (서울공용2로 변경)

start_row = 4
for i, row_data in enumerate(data):
    r = start_row + i
    ws.row_dimensions[r].height = 18
    is_changed = (i == changed_row_idx)

    for j, val in enumerate(row_data):
        col = j + 1
        cell = ws.cell(row=r, column=col)
        cell.value = val
        cell.border = border
        cell.font = black_font

        if is_changed:
            cell.fill = changed_fill

        # 숫자 서식
        if col in [9, 10, 11, 12]:
            cell.alignment = right_align
            if isinstance(val, int):
                cell.number_format = '#,##0'
        elif col in [6, 7]:
            cell.alignment = center
        else:
            cell.alignment = center

# ── 합계 행 ───────────────────────────────────────────
total_row = start_row + len(data)
ws.row_dimensions[total_row].height = 20
ws.merge_cells(f'A{total_row}:H{total_row}')
tc = ws[f'A{total_row}']
tc.value = '합계'
tc.font = total_font
tc.fill = total_fill
tc.alignment = center
tc.border = border

# 합계 계산 (200호2725 제외 후)
sum_deposit = sum(d[8] for d in data if d[8] is not None)
sum_rent = sum(d[9] for d in data if d[9] is not None)
sum_etc = sum(d[10] for d in data if d[10] is not None)
sum_total = sum(d[11] for d in data if d[11] is not None)

for col, val in [(9, sum_deposit), (10, sum_rent), (11, sum_etc), (12, sum_total)]:
    cell = ws.cell(row=total_row, column=col)
    cell.value = val
    cell.font = total_font
    cell.fill = total_fill
    cell.alignment = right_align
    cell.border = border
    cell.number_format = '#,##0'

ws.cell(row=total_row, column=13).border = border
ws.cell(row=total_row, column=13).fill = total_fill

# 비고 행 (변경사항 안내)
note_row = total_row + 2
ws.merge_cells(f'A{note_row}:M{note_row}')
note_cell = ws[f'A{note_row}']
note_cell.value = '※ 변경사항: ① 200호2725 (오진택) 계약해지 → 목록에서 제외  ② 200호3807 → 경영기획실 공용차량(서울공용2)으로 변경 (노란색 표시)'
note_cell.font = Font(name='맑은 고딕', size=9, color='C00000', bold=True)
note_cell.alignment = Alignment(horizontal='left', vertical='center')

ws.row_dimensions[note_row].height = 18

# 열 고정 (헤더 고정)
ws.freeze_panes = 'A4'

output_path = '/home/user/NRB/차량_렌탈현황_변경.xlsx'
wb.save(output_path)
print(f"저장 완료: {output_path}")
print(f"총 차량 수: {len(data)}대 (200호2725 계약해지로 1대 제외)")
print(f"합계 - 보증금: {sum_deposit:,}  월렌트료: {sum_rent:,}  기타: {sum_etc:,}  총계: {sum_total:,}")
