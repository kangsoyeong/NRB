import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ── 5회차 실제 공장/현장 금액 (이미지 기준) ─────────────────────
FACTORY_BASE = 2_347_000_000
FIELD_BASE   = 1_633_000_000
TOTAL_BASE   = 3_980_000_000
FACTORY_PCT  = FACTORY_BASE / TOTAL_BASE   # 0.589698...
FIELD_PCT    = FIELD_BASE   / TOTAL_BASE   # 0.410302...

def split(total, actual_factory=None):
    if actual_factory is not None:
        return actual_factory, total - actual_factory
    f = round(total * FACTORY_BASE / TOTAL_BASE)
    return f, total - f

# ── 스타일 헬퍼 ─────────────────────────────────────────────────
thin   = Side(style="thin")
med    = Side(style="medium")

def bdr(l=thin, r=thin, t=thin, b=thin):
    return Border(left=l, right=r, top=t, bottom=b)

def fl(hex_): return PatternFill("solid", fgColor=hex_)
def ft(bold=False, color="000000", size=10, italic=False):
    return Font(name="맑은 고딕", bold=bold, color=color, size=size, italic=italic)
def al(h="center", v="center", wrap=True):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

NUM  = '#,##0'
PCT  = '0.00%'

# 색상
T_DARK  = "1F4E79"   # 제목
H1      = "2E75B6"   # 메인 헤더
H2      = "BDD7EE"   # 서브 헤더
H_FAC   = "D6E4F0"   # 공장 열 헤더
H_FLD   = "FCE4D6"   # 현장 열 헤더
H_TOT   = "E2EFDA"   # 합계 열 헤더
ROW_FAC = "EBF3FA"   # 공장 데이터 행
ROW_FLD = "FFF2E8"   # 현장 데이터 행
ROW_TOT = "F2F9F0"   # 합계 데이터 행
ROW_ALT = "F5F5F5"
WHITE   = "FFFFFF"
YELLOW  = "FFF2CC"

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "도급기성청구현황"

# ── 열 너비 ──────────────────────────────────────────────────────
col_w = {
    "A": 7,   # 회차
    "B": 7,   # 구분(정식/약식)
    "C": 20,  # 품의번호
    "D": 12,  # 기안일자
    "E": 12,  # 처리일자
    "F": 14,  # 청구월
    "G": 18,  # 공장분 금액
    "H": 9,   # 공장 비율
    "I": 18,  # 현장분 금액
    "J": 9,   # 현장 비율
    "K": 18,  # 금회합계
    "L": 17,  # 극동건설
    "M": 15,  # HJ중공업
    "N": 15,  # 서한
    "O": 15,  # 대저건설
    "P": 15,  # 엔알비
    "Q": 20,  # 비고
}
for col, w in col_w.items():
    ws.column_dimensions[col].width = w

# ═══ 1행: 제목 ═══════════════════════════════════════════════════
ws.merge_cells("A1:Q1")
c = ws["A1"]
c.value = "의왕초평 A-4BL  도급 기성청구 현황 (공장/현장 분리)"
c.font      = Font(name="맑은 고딕", bold=True, size=15, color="FFFFFF")
c.alignment = al()
c.fill      = fl(T_DARK)
c.border    = bdr()
ws.row_dimensions[1].height = 38

# ═══ 2행: 비율 기준 안내 ════════════════════════════════════════
ws.merge_cells("A2:F2")
ws["A2"].value     = "【공장/현장 분리 기준 비율】"
ws["A2"].font      = ft(bold=True, size=10)
ws["A2"].alignment = al(h="left")
ws["A2"].fill      = fl(H2); ws["A2"].border = bdr()

ws.merge_cells("G2:H2")
ws["G2"].value         = f"공장(제1사업장): {FACTORY_BASE:,}원"
ws["G2"].font          = ft(bold=True, color="1F4E79", size=10)
ws["G2"].alignment     = al(h="center")
ws["G2"].fill          = fl(H_FAC); ws["G2"].border = bdr()

ws.merge_cells("I2:J2")
ws["I2"].value         = f"현장(제2사업장): {FIELD_BASE:,}원"
ws["I2"].font          = ft(bold=True, color="843C0C", size=10)
ws["I2"].alignment     = al(h="center")
ws["I2"].fill          = fl(H_FLD); ws["I2"].border = bdr()

ws.merge_cells("K2:K2")
ws["K2"].value     = f"기준합계: {TOTAL_BASE:,}원"
ws["K2"].font      = ft(bold=True, size=9)
ws["K2"].alignment = al()
ws["K2"].fill      = fl(H_TOT); ws["K2"].border = bdr()

ws.merge_cells("L2:Q2")
ws["L2"].value     = "※ 5회차 실제 데이터 기준 비율 / 1~4회차는 동일 비율 적용 (추정)"
ws["L2"].font      = ft(italic=True, size=9, color="595959")
ws["L2"].alignment = al(h="left")
ws["L2"].fill      = fl(H2); ws["L2"].border = bdr()
ws.row_dimensions[2].height = 22

# ═══ 3행: 대분류 헤더 ═══════════════════════════════════════════
HEADER_MAP = {  # col_letter: (merge_end, label, fill_color)
    "A": ("B",  "기본정보",           H1),
    "C": ("F",  "기본정보",           H1),
    "G": ("H",  f"공장(제1사업장)\n{FACTORY_PCT*100:.2f}%",  H_FAC),
    "I": ("J",  f"현장(제2사업장)\n{FIELD_PCT*100:.2f}%",    H_FLD),
    "K": ("K",  "금회합계",           H_TOT),
    "L": ("P",  "업체별 금회금액",    H1),
    "Q": ("Q",  "비고",               H1),
}
skip_cols = set()
for start_col, (end_col, label, color) in HEADER_MAP.items():
    if start_col == end_col:
        c = ws[f"{start_col}3"]
        c.value = label
    else:
        ws.merge_cells(f"{start_col}3:{end_col}3")
        c = ws[f"{start_col}3"]
        c.value = label
        for col in range(
            openpyxl.utils.column_index_from_string(start_col)+1,
            openpyxl.utils.column_index_from_string(end_col)+1
        ):
            skip_cols.add(col)
    text_color = "FFFFFF" if color in (H1, T_DARK) else ("1F4E79" if color == H_FAC else ("843C0C" if color == H_FLD else "375623"))
    c.font      = Font(name="맑은 고딕", bold=True, size=10, color=text_color)
    c.alignment = al()
    c.fill      = fl(color)
    c.border    = bdr()
ws.row_dimensions[3].height = 32

# ═══ 4행: 소분류 헤더 ═══════════════════════════════════════════
sub_headers = [
    ("A", "회차"),
    ("B", "정식\n/약식"),
    ("C", "품의번호"),
    ("D", "기안일자"),
    ("E", "처리일자"),
    ("F", "청구월"),
    ("G", "공장분\n금액"),
    ("H", "공장\n비율"),
    ("I", "현장분\n금액"),
    ("J", "현장\n비율"),
    ("K", "합 계"),
    ("L", "극동건설(주)\n51%"),
    ("M", "(주)HJ중공업\n15%"),
    ("N", "(주)서한\n14%"),
    ("O", "(주)대저건설\n10%"),
    ("P", "(주)엔알비\n10%"),
    ("Q", "비고"),
]
fac_cols = {"G", "H"}
fld_cols = {"I", "J"}
tot_cols = {"K"}
for col_letter, label in sub_headers:
    c = ws[f"{col_letter}4"]
    c.value = label
    if col_letter in fac_cols:
        bg, tc = H_FAC, "1F4E79"
    elif col_letter in fld_cols:
        bg, tc = H_FLD, "843C0C"
    elif col_letter in tot_cols:
        bg, tc = H_TOT, "375623"
    else:
        bg, tc = H1, "FFFFFF"
    c.font      = Font(name="맑은 고딕", bold=True, size=9, color=tc)
    c.alignment = al()
    c.fill      = fl(bg)
    c.border    = bdr()
ws.row_dimensions[4].height = 36

# ═══ 데이터 ════════════════════════════════════════════════════
rows_data = [
    # (회차, 구분, 품의번호, 기안일자, 처리일자, 청구월,
    #  합계, 극동, HJ, 서한, 대저, 엔알비, actual_factory or None, 비고)
    ("1회", "정식", "(주)엔알비-2025-17478", "2025-11-29", "2025-12-04", "2025년 12월",
     5_076_000_000, 2_588_760_000, 761_400_000, 710_640_000, 507_600_000, 507_600_000, None,
     "기안일(11/29) vs 처리일(12/04) 월 상이"),
    ("2회", "약식", "(주)엔알비-2026-1859",  "2026-02-02", "2026-02-04", "2026년 2월",
     624_000_000, 318_240_000, 93_600_000, 87_360_000, 62_400_000, 62_400_000, None, ""),
    ("3회", "약식", "(주)엔알비-2026-4026",  "2026-03-05", "2026-03-06", "2026년 3월",
     1_640_000_000, 836_400_000, 246_000_000, 229_600_000, 164_000_000, 164_000_000, None, ""),
    ("4회", "정식", "(주)엔알비-2026-5251",  "2026-03-25", "2026-04-06", "2026년 4월",
     2_930_000_000, 1_494_300_000, 439_500_000, 410_200_000, 293_000_000, 293_000_000, None,
     "기안일(3/25) vs 처리일(4/06) 월 상이"),
    ("5회", "약식", "(주)엔알비-2026-7755",  "미확인",     "2026-05-12", "2026년 5월",
     3_980_000_000, 2_029_800_000, 597_000_000, 557_200_000, 398_000_000, 398_000_000, FACTORY_BASE,
     "5회차 실제값 기준"),
]

WARN = {0, 3}   # 1회·4회 (월 상이)

for i, rd in enumerate(rows_data):
    row = 5 + i
    (회차, 구분, 품의, 기안, 처리, 청구월,
     합계, 극동, hj, 서한, 대저, 엔알비, act_fac, 비고) = rd

    fac_amt, fld_amt = split(합계, act_fac)
    fac_r = fac_amt / 합계
    fld_r = fld_amt / 합계

    row_bg = YELLOW if i in WARN else (ROW_ALT if i % 2 else WHITE)

    values = [
        (회차,       ft(bold=True, size=11),    al(),         row_bg, None),
        (구분,       ft(bold=True, size=9,
                        color=("C00000" if 구분=="정식" else "2E75B6")),
                     al(),         row_bg, None),
        (품의,       ft(size=9),                al(h="left"), row_bg, None),
        (기안,       ft(size=9),                al(),         row_bg, None),
        (처리,       ft(size=9),                al(),         row_bg, None),
        (청구월,     ft(bold=True, size=10),    al(),         row_bg, None),
        (fac_amt,    ft(bold=True, color="1F4E79", size=10),
                     al(h="right"), ROW_FAC,   NUM),
        (fac_r,      ft(bold=True, color="1F4E79", size=10),
                     al(),          ROW_FAC,   PCT),
        (fld_amt,    ft(bold=True, color="843C0C", size=10),
                     al(h="right"), ROW_FLD,   NUM),
        (fld_r,      ft(bold=True, color="843C0C", size=10),
                     al(),          ROW_FLD,   PCT),
        (합계,       ft(bold=True, size=10),    al(h="right"), ROW_TOT, NUM),
        (극동,       ft(size=9),               al(h="right"), row_bg, NUM),
        (hj,         ft(size=9),               al(h="right"), row_bg, NUM),
        (서한,       ft(size=9),               al(h="right"), row_bg, NUM),
        (대저,       ft(size=9),               al(h="right"), row_bg, NUM),
        (엔알비,     ft(size=9),               al(h="right"), row_bg, NUM),
        (비고,       ft(italic=True, size=8,
                        color="595959"),       al(h="left"), row_bg, None),
    ]

    for col_idx, (val, fnt, aln, bg, nfmt) in enumerate(values, start=1):
        c = ws.cell(row=row, column=col_idx, value=val)
        c.font   = fnt
        c.alignment = aln
        c.fill   = fl(bg)
        c.border = bdr()
        if nfmt:
            c.number_format = nfmt

    ws.row_dimensions[row].height = 26

# ═══ 합계 행 ════════════════════════════════════════════════════
total_row = 5 + len(rows_data)
ws.merge_cells(f"A{total_row}:F{total_row}")
c = ws[f"A{total_row}"]
c.value     = "합   계"
c.font      = Font(name="맑은 고딕", bold=True, size=11, color="FFFFFF")
c.alignment = al(); c.fill = fl(T_DARK); c.border = bdr()

total_amt = sum(rd[6] for rd in rows_data)
total_fac = sum(split(rd[6], rd[12])[0] for rd in rows_data)
total_fld = total_amt - total_fac

total_vals = [
    ("G", total_fac,             NUM,  "1F4E79"),
    ("H", total_fac/total_amt,   PCT,  "1F4E79"),
    ("I", total_fld,             NUM,  "843C0C"),
    ("J", total_fld/total_amt,   PCT,  "843C0C"),
    ("K", total_amt,             NUM,  "FFFFFF"),
    ("L", sum(rd[7]  for rd in rows_data), NUM, "FFFFFF"),
    ("M", sum(rd[8]  for rd in rows_data), NUM, "FFFFFF"),
    ("N", sum(rd[9]  for rd in rows_data), NUM, "FFFFFF"),
    ("O", sum(rd[10] for rd in rows_data), NUM, "FFFFFF"),
    ("P", sum(rd[11] for rd in rows_data), NUM, "FFFFFF"),
    ("Q", "",                     None, "FFFFFF"),
]
for col_letter, val, nfmt, tc in total_vals:
    c = ws[f"{col_letter}{total_row}"]
    c.value = val
    c.font  = Font(name="맑은 고딕", bold=True, size=10, color=tc)
    c.alignment = al(h="right")
    bg_map = {"G":"1F4E79","H":"1F4E79","I":"843C0C","J":"843C0C","K":T_DARK}
    c.fill  = fl(bg_map.get(col_letter, T_DARK))
    c.border = bdr()
    if nfmt: c.number_format = nfmt
ws.row_dimensions[total_row].height = 28

# ═══ 각주 ════════════════════════════════════════════════════════
note_r = total_row + 2
ws.merge_cells(f"A{note_r}:Q{note_r}")
c = ws[f"A{note_r}"]
c.value = (
    f"※ 공장(제1사업장) {FACTORY_PCT*100:.2f}% / 현장(제2사업장) {FIELD_PCT*100:.2f}%  "
    f"← 5회차 실제값(공장 {FACTORY_BASE:,}원 / 현장 {FIELD_BASE:,}원 / 합계 {TOTAL_BASE:,}원) 기준 산정. "
    "1~4회차는 동일 비율 추정 적용."
)
c.font      = Font(name="맑은 고딕", italic=True, size=9, color="C00000")
c.alignment = al(h="left")

# ═══ 인쇄 ════════════════════════════════════════════════════════
ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
ws.page_setup.fitToPage   = True
ws.page_setup.fitToWidth  = 1
ws.freeze_panes           = "G5"

out = "/home/user/NRB/의왕초평_도급기성청구현황.xlsx"
wb.save(out)
print(f"저장 완료: {out}")
print(f"\n비율 요약: 공장 {FACTORY_PCT*100:.2f}% / 현장 {FIELD_PCT*100:.2f}%")
print(f"합계: 공장 {total_fac:,}  현장 {total_fld:,}  총 {total_amt:,}")
