import openpyxl
from openpyxl.styles import (
    Font, Alignment, PatternFill, Border, Side, numbers
)
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "차입금 현황 요약"

# ── 색상 정의 ──
NAVY   = "1F3864"
BLUE   = "2E5EA3"
LBLUE  = "D9E1F2"
LGRAY  = "F2F2F2"
WHITE  = "FFFFFF"
YELLOW = "FFFF00"
RED    = "FF0000"
ORANGE = "FCE4D6"

def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def font(bold=False, size=9, color="000000", name="맑은 고딕"):
    return Font(name=name, size=size, bold=bold, color=color)

def align(h="center", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def border(style="thin", color="BFBFBF"):
    s = Side(style=style, color=color)
    return Border(left=s, right=s, top=s, bottom=s)

def thick_border():
    tk = Side(style="medium", color="000000")
    tn = Side(style="thin", color="BFBFBF")
    return Border(left=tk, right=tk, top=tk, bottom=tk)

def apply(cell, value=None, bold=False, size=9, fc=None, tc="000000",
          h="center", v="center", wrap=False, num_fmt=None, bdr=None):
    if value is not None:
        cell.value = value
    cell.font      = font(bold=bold, size=size, color=tc)
    cell.alignment = align(h=h, v=v, wrap=wrap)
    if fc:
        cell.fill = fill(fc)
    if bdr:
        cell.border = bdr
    if num_fmt:
        cell.number_format = num_fmt

# ── 열 너비 설정 ──
col_widths = {
    "A": 8,   # 유형
    "B": 10,  # 구분
    "C": 16,  # 거래처명
    "D": 7,   # 건수
    "E": 18,  # 최초 차입액
    "F": 18,  # 금월 잔액
    "G": 13,  # 가중평균금리
}
for col, w in col_widths.items():
    ws.column_dimensions[col].width = w

# ── 행 높이 ──
for r in range(1, 25):
    ws.row_dimensions[r].height = 18

# ── 제목 행 ──
ws.merge_cells("A1:G1")
c = ws["A1"]
apply(c, "NRB 차입금 현황 요약표", bold=True, size=12, fc=NAVY, tc="FFFFFF")
ws.row_dimensions[1].height = 28

# ── 기준일 ──
ws.merge_cells("A2:G2")
apply(ws["A2"], "기준일: 2026년 5월 31일", bold=False, size=9, h="right", fc=LGRAY)
ws.row_dimensions[2].height = 16

# ── 헤더 행 ──
headers = ["유형", "구분", "거래처명", "건수", "최초 차입액", "금월 잔액", "가중평균금리"]
for i, h_val in enumerate(headers, start=1):
    c = ws.cell(row=3, column=i)
    apply(c, h_val, bold=True, size=9, fc=BLUE, tc="FFFFFF", bdr=border("thin","FFFFFF"))
ws.row_dimensions[3].height = 22

# ── 데이터 정의 ──
# (유형, 구분, 거래처명, 건수, 최초차입액, 금월잔액, 가중평균금리, row_type)
# row_type: "data" / "sub" / "total" / "note"
data_rows = [
    # 은행 - 기업
    ("차입금", "은행", "기업은행",    22, 59_270_000_000, 54_829_098_254, 0.037900, "data"),
    ("",       "",    "광주은행",      1,  7_000_000_000,    874_998_000, 0.054900, "data"),
    ("",       "",    "전북은행",      2,  9_000_000_000,  2_376_333_331, 0.050500, "data"),
    ("",       "",    "신협은행",      1,  2_000_000_000,    694_363_893, 0.063000, "data"),
    ("",       "",    "우리은행",      2, 12_000_000_000, 10_500_000_000, 0.038371, "data"),
    ("",       "",    "소기업진흥공단", 1,  5_000_000_000,    416_400_000, 0.026700, "data_note"),
    ("",       "소계","",            29, 94_270_000_000, 69_691_193_478, None,      "sub"),
    # 전환사채
    ("",  "전환사채", "에스엘아이 퀀텀 성장 2호 펀드", 1, 5_000_004_800, 5_000_004_800, 0.020000, "data"),
    ("",  "",        "우리 2022 스케일업 펀드",        1, 5_000_004_800, 5_000_004_800, 0.020000, "data"),
    ("",  "소계",    "",                               2, 10_000_009_600, 10_000_009_600, None, "sub"),
    # 회사채
    ("",  "회사채",  "수협 2025 스케일업",             1,  3_000_000_000,  3_000_000_000, 0.060280, "data"),
    ("",  "소계",    "",                               1,  3_000_000_000,  3_000_000_000, None,     "sub"),
    # 합계
    ("합계", "", "",  32, 107_270_009_600, 82_691_203_078, 0.037300, "total"),
]

# ── 셀 작성 ──
START_ROW = 4

# 유형 병합을 위한 추적
merge_ranges = {
    "유형_차입금": (START_ROW, START_ROW + len(data_rows) - 2),  # 합계 행 제외
}

for row_idx, (유형, 구분, 거래처, 건수, 최초, 잔액, 금리, rtype) in enumerate(data_rows):
    r = START_ROW + row_idx

    if rtype == "total":
        fc_row = BLUE
        tc_row = "FFFFFF"
        bold_row = True
    elif rtype == "sub":
        fc_row = LBLUE
        tc_row = "000000"
        bold_row = True
    elif rtype == "data_note":
        fc_row = ORANGE
        tc_row = "000000"
        bold_row = False
    else:
        fc_row = WHITE
        tc_row = "000000"
        bold_row = False

    bdr = border()

    # A: 유형
    apply(ws.cell(r, 1), 유형 or None, bold=bold_row, size=9,
          fc=fc_row, tc=tc_row, bdr=bdr)
    # B: 구분
    apply(ws.cell(r, 2), 구분 or None, bold=bold_row, size=9,
          fc=fc_row, tc=tc_row, bdr=bdr)
    # C: 거래처
    apply(ws.cell(r, 3), 거래처 or None, bold=bold_row, size=9,
          fc=fc_row, tc=tc_row, h="left", bdr=bdr)
    # D: 건수
    apply(ws.cell(r, 4), 건수, bold=bold_row, size=9,
          fc=fc_row, tc=tc_row, num_fmt="#,##0", bdr=bdr)
    # E: 최초
    apply(ws.cell(r, 5), 최초, bold=bold_row, size=9,
          fc=fc_row, tc=tc_row, num_fmt="#,##0", bdr=bdr)
    # F: 잔액
    apply(ws.cell(r, 6), 잔액, bold=bold_row, size=9,
          fc=fc_row, tc=tc_row, num_fmt="#,##0", bdr=bdr)
    # G: 금리
    if 금리 is not None:
        apply(ws.cell(r, 7), 금리, bold=bold_row, size=9,
              fc=fc_row, tc=tc_row, num_fmt="0.0000%", bdr=bdr)
    else:
        apply(ws.cell(r, 7), "-", bold=bold_row, size=9,
              fc=fc_row, tc=tc_row, bdr=bdr)

# ── 유형 열(A) 병합: 차입금 ──
ws.merge_cells(f"A{START_ROW}:A{START_ROW + len(data_rows) - 2}")
c = ws[f"A{START_ROW}"]
apply(c, "차입금", bold=True, size=10, fc=LGRAY, tc="000000")

# ── 구분 열(B) 병합: 은행 ──
ws.merge_cells(f"B{START_ROW}:B{START_ROW + 6}")
c = ws[f"B{START_ROW}"]
apply(c, "은행", bold=True, size=9, fc=LGRAY, tc="000000")

# ── 주석 행 ──
note_r = START_ROW + len(data_rows)
ws.row_dimensions[note_r].height = 14
ws.merge_cells(f"A{note_r}:G{note_r}")
apply(ws.cell(note_r, 1),
      "※ 건수 합계는 32건(은행 29 + 전환사채 2 + 회사채 1)이 맞으며, 수입신용장(기업은행 US003) 잔액은 기업은행 금월잔액에 포함됨",
      size=8, tc="FF0000", h="left", fc="FFFCE4")

note_r2 = note_r + 1
ws.row_dimensions[note_r2].height = 14
ws.merge_cells(f"A{note_r2}:G{note_r2}")
apply(ws.cell(note_r2, 1),
      "※ 소기업진흥공단 항목(주황색)은 제공된 원본 데이터에 미포함 — 별도 확인 필요 (최초차입금 5,000,000,000 / 잔액 416,400,000 / 금리 2.67%)",
      size=8, tc="C55A11", h="left", fc="FFFCE4")

wb.save("/home/user/NRB/차입금현황_요약표_검토본.xlsx")
print("완료")
