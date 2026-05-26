import openpyxl

FILE1 = "/root/.claude/uploads/24041d56-69af-4927-919e-4cbd89e59ed0/35a33838-_____260523.xlsx"
FILE2 = "/root/.claude/uploads/24041d56-69af-4927-919e-4cbd89e59ed0/875360de-__________.xlsx"

wb1 = openpyxl.load_workbook(FILE1, data_only=True)
ws1 = wb1["Cash out"]

wb2 = openpyxl.load_workbook(FILE2)
ws2 = wb2["Sheet1"]

COL1_DOC = 1   # B
COL1_SRC = 16  # Q
COL2_DOC = 14  # O
COL2_SRC = 15  # P

# Build mapping
mapping = {}
for row in ws1.iter_rows(min_row=3, values_only=True):
    doc = row[COL1_DOC] if len(row) > COL1_DOC else None
    src = row[COL1_SRC] if len(row) > COL1_SRC else None
    if doc is None or str(doc).strip() == "":
        continue
    if src is None or str(src).strip() == "":
        continue
    doc_key = str(doc).strip()
    src_val = str(src).strip()
    if doc_key not in mapping:
        mapping[doc_key] = []
    if src_val not in mapping[doc_key]:
        mapping[doc_key].append(src_val)

print(f"Mapping has {len(mapping)} entries")
print("\nSample File1 keys (first 10):")
for k in list(mapping.keys())[:10]:
    print(f"  repr={repr(k)}")

# Collect blank rows from File2
print("\nBlank rows in File2 (col O = 지결문서번호, col P = 원천전표번호):")
blank_rows = []
for i, row_cells in enumerate(ws2.iter_rows(min_row=2), start=2):
    doc_cell = row_cells[COL2_DOC]
    src_cell = row_cells[COL2_SRC]
    src_val = src_cell.value
    src_blank = (src_val is None or str(src_val).strip() == "")
    if src_blank:
        doc_val = doc_cell.value
        blank_rows.append((i, doc_val, doc_cell))

print(f"  Total blank: {len(blank_rows)}")
print("\nBlank row details (row, doc_value, repr):")
for row_num, doc_val, doc_cell in blank_rows[:20]:
    in_map = str(doc_val).strip() in mapping if doc_val else False
    print(f"  row {row_num:5d}  doc={repr(doc_val)}  in_map={in_map}")

# Also check type info
print("\nType info for first few blank rows:")
for row_num, doc_val, doc_cell in blank_rows[:5]:
    print(f"  row {row_num}  type={type(doc_val).__name__}  value={repr(doc_val)}")

# Show some mapping keys to compare
print("\nFirst 10 mapping keys repr:")
for k in list(mapping.keys())[:10]:
    print(f"  {repr(k)}")
