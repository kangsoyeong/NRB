import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import copy

PATH = '/home/user/NRB/의왕초평_도급기성청구현황.xlsx'

wb = openpyxl.load_workbook(PATH)
ws = wb.active

# ── 스타일 헬퍼 ─────────────────────────────────────────────────
thin = Side(style="thin")
med  = Side(style="medium")

def bdr(all_thin=True):
    s = thin if all_thin else med
    return Border(left=s, right=s, top=s, bottom=s)

def fl(hex_): return PatternFill("solid", fgColor=hex_)
def ft(bold=False, color="000000", size=10, italic=False):
    return Font(name="맑은 고딕", bold=bold, color=color, size=size, italic=italic)
def al(h="center", v="center", wrap=True):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

NUM = '#,##0'

C_DARK  = "1F4E79"
C_MID   = "2E75B6"
C_M4    = "E2EFDA"  # 4월 (연초록)
C_M5    = "FCE4D6"  # 5월 (연주황)
C_SUM   = "BDD7EE"  # 입금합계
C_UNP   = "FFE0E0"  # 미수금 (연빨강)
C_ALT   = "F5F5F5"
C_WHITE = "FFFFFF"
C_HEAD  = "D9E1F2"

# ── 4회차 입금 데이터 ────────────────────────────────────────────
# (업체명, 청구금액, 4월입금, 4월일자, 5월입금, 5월일자, 비고)
payments = [
    ("극동건설(주)",  1_494_300_000, None,        None,  1_494_300_000, "5/7",  "완납"),
    ("(주)HJ중공업",   439_500_000, None,        None,          None,   None,  "미수금"),
    ("(주)서한",       410_200_000, None,        None,    410_200_000, "5/7",  "완납"),
    ("(주)대저건설",   293_000_000, 293_000_000, "4/24",        None,   None,  "완납"),
    ("(주)엔알비(LH)", 293_000_000, 293_000_000, "4/23",        None,   None,  "완납"),
]

# ── 시작 행: 기존 마지막 행 + 3 ──────────────────────────────────
start = ws.max_row + 3

# ── 섹션 제목 ────────────────────────────────────────────────────
ws.merge_cells(f"A{start}:I{start}")
c = ws[f"A{start}"]
c.value     = "【 4회차 입금 현황 】"
c.font      = Font(name="맑은 고딕", bold=True, size=13, color="FFFFFF")
c.alignment = al()
c.fill      = fl(C_DARK)
c.border    = bdr()
ws.row_dimensions[start].height = 30

# ── 헤더 행 ─────────────────────────────────────────────────────
hr = start + 1
hdr_data = [
    ("A", "업체명"),
    ("B", "청구금액"),
    ("C", "4월 입금"),
    ("D", "4월\n입금일"),
    ("E", "5월 입금"),
    ("F", "5월\n입금일"),
    ("G", "입금합계"),
    ("H", "미수금"),
    ("I", "비고"),
]
col_fills = {
    "A": C_HEAD, "B": C_HEAD,
    "C": C_M4,   "D": C_M4,
    "E": C_M5,   "F": C_M5,
    "G": C_SUM,
    "H": C_UNP,
    "I": C_HEAD,
}
for col_letter, label in hdr_data:
    c = ws[f"{col_letter}{hr}"]
    c.value     = label
    c.font      = ft(bold=True, size=9,
                     color=("FFFFFF" if col_fills[col_letter] == C_DARK else "000000"))
    c.alignment = al()
    c.fill      = fl(col_fills[col_letter])
    c.border    = bdr()
ws.row_dimensions[hr].height = 34

# ── 데이터 행 ────────────────────────────────────────────────────
for i, (업체, 청구, m4, d4, m5, d5, note) in enumerate(payments):
    row  = hr + 1 + i
    paid = (m4 or 0) + (m5 or 0)
    unp  = 청구 - paid
    bg   = C_UNP if unp > 0 else (C_ALT if i % 2 == 0 else C_WHITE)

    row_vals = [
        ("A", 업체,   ft(bold=True, size=9),    al(h="left"), bg,   None),
        ("B", 청구,   ft(size=9),               al(h="right"), bg,  NUM),
        ("C", m4,     ft(bold=True if m4 else False, size=9,
                         color="375623" if m4 else "AAAAAA"),
                      al(h="right"), C_M4 if m4 else fl(C_M4) and C_M4, NUM),
        ("D", d4 or "-", ft(size=9),            al(),         C_M4, None),
        ("E", m5,     ft(bold=True if m5 else False, size=9,
                         color="843C0C" if m5 else "AAAAAA"),
                      al(h="right"), C_M5, NUM),
        ("F", d5 or "-", ft(size=9),            al(),         C_M5, None),
        ("G", paid if paid else None,
              ft(bold=True, size=9, color="1F4E79" if paid else "AAAAAA"),
              al(h="right"), C_SUM, NUM),
        ("H", unp if unp > 0 else None,
              ft(bold=True, size=9, color="C00000" if unp > 0 else "000000"),
              al(h="right"), C_UNP, NUM),
        ("I", note,   ft(italic=True, size=8,
                         color=("C00000" if note == "미수금" else "375623")),
                      al(), bg, None),
    ]

    for col_letter, val, fnt, aln, bg_c, nfmt in row_vals:
        c = ws[f"{col_letter}{row}"]
        c.value     = val
        c.font      = fnt
        c.alignment = aln
        c.fill      = fl(bg_c) if isinstance(bg_c, str) else fl(bg)
        c.border    = bdr()
        if nfmt and val is not None:
            c.number_format = nfmt

    ws.row_dimensions[row].height = 22

# ── 합계 행 ─────────────────────────────────────────────────────
total_row = hr + 1 + len(payments)
total_bill = sum(p[1] for p in payments)
total_m4   = sum(p[2] or 0 for p in payments)
total_m5   = sum(p[4] or 0 for p in payments)
total_paid = total_m4 + total_m5
total_unp  = total_bill - total_paid

ws.merge_cells(f"A{total_row}:A{total_row}")
for col_letter, val, nfmt, bg in [
    ("A", "합   계", None, C_DARK),
    ("B", total_bill, NUM, C_DARK),
    ("C", total_m4,   NUM, C_M4),
    ("D", "—",        None, C_M4),
    ("E", total_m5,   NUM, C_M5),
    ("F", "—",        None, C_M5),
    ("G", total_paid, NUM, C_SUM),
    ("H", total_unp if total_unp > 0 else None, NUM, C_UNP),
    ("I", "",         None, C_DARK),
]:
    c = ws[f"{col_letter}{total_row}"]
    c.value     = val
    c.font      = Font(name="맑은 고딕", bold=True, size=10,
                       color="FFFFFF" if bg in (C_DARK,) else
                             ("375623" if bg == C_M4 else
                              ("843C0C" if bg == C_M5 else
                               ("1F4E79" if bg == C_SUM else
                                ("C00000" if bg == C_UNP and val else "000000")))))
    c.alignment = al(h="right" if isinstance(val, int) else "center")
    c.fill      = fl(bg)
    c.border    = bdr()
    if nfmt and val:
        c.number_format = nfmt
ws.row_dimensions[total_row].height = 24

# ── 월별 요약 박스 ────────────────────────────────────────────────
sr = total_row + 2
ws.merge_cells(f"A{sr}:I{sr}")
ws[f"A{sr}"].value     = "▶ 월별 입금 요약"
ws[f"A{sr}"].font      = ft(bold=True, size=10, color="FFFFFF")
ws[f"A{sr}"].alignment = al(h="left")
ws[f"A{sr}"].fill      = fl(C_MID)
ws[f"A{sr}"].border    = bdr()
ws.row_dimensions[sr].height = 22

summary = [
    ("4월", total_m4, "대저건설(4/24 293,000,000) + 엔알비(4/23 293,000,000)"),
    ("5월", total_m5, "극동건설(5/7 1,494,300,000) + 서한(5/7 410,200,000)"),
    ("미수금", total_unp, "(주)HJ중공업 미입금"),
]
for j, (label, amt, detail) in enumerate(summary):
    row = sr + 1 + j
    is_unp = label == "미수금"
    bg = C_UNP if is_unp else (C_M4 if label == "4월" else C_M5)

    ws.merge_cells(f"A{row}:A{row}")
    c = ws[f"A{row}"]
    c.value     = label
    c.font      = ft(bold=True, size=10,
                     color="C00000" if is_unp else ("375623" if label=="4월" else "843C0C"))
    c.alignment = al()
    c.fill      = fl(bg)
    c.border    = bdr()

    ws.merge_cells(f"B{row}:C{row}")
    c = ws[f"B{row}"]
    c.value         = amt
    c.font          = ft(bold=True, size=10,
                         color="C00000" if is_unp else "1F4E79")
    c.alignment     = al(h="right")
    c.fill          = fl(bg)
    c.border        = bdr()
    c.number_format = NUM

    ws.merge_cells(f"D{row}:I{row}")
    c = ws[f"D{row}"]
    c.value     = detail
    c.font      = ft(italic=True, size=9, color="595959")
    c.alignment = al(h="left")
    c.fill      = fl(C_WHITE)
    c.border    = bdr()

    ws.row_dimensions[row].height = 22

wb.save(PATH)
print("저장 완료:", PATH)
print(f"\n4회차 입금 요약:")
print(f"  4월 입금: {total_m4:>15,}원")
print(f"  5월 입금: {total_m5:>15,}원")
print(f"  입금합계: {total_paid:>15,}원")
print(f"  미수금:   {total_unp:>15,}원 (HJ중공업)")
print(f"  청구합계: {total_bill:>15,}원")
