# -*- coding: utf-8 -*-
"""
차입금 이자/원금 상환 스케줄 생성
대상: 전북은행 / 신협은행 / 광주은행
기준: '1-1. 차입금현황' (2026년 5월말 잔액 기준 → 6월부터 만기까지 예상 스케줄)
"""
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DAY_BASIS = 365

# ---------------------------------------------------------------------------
# 차입금 정보 (2026-05-31 월말 잔액 기준)
# ---------------------------------------------------------------------------
LOANS = [
    {
        "bank": "전북은행",
        "name": "기운인대(원금균등)-6831557",
        "account": "1071-02-6831557",
        "pay_day": 30,
        "rate": 0.0549,
        "maturity": datetime.date(2026, 12, 30),
        "method": "원금균등분할상환",
        "monthly_principal": 208333333,   # 매월 균등 원금
        "balance_0531": 1458333331,       # 2026-05-31 잔액
    },
    {
        "bank": "전북은행",
        "name": "JB관계형금융우대론(3781411)",
        "account": "1071-02-3781411",
        "pay_day": 27,
        "rate": 0.0436,
        "maturity": datetime.date(2026, 12, 27),
        "method": "자유분할상환(원금)",
        "monthly_principal": 134000000,
        "balance_0531": 918000000,
    },
    {
        "bank": "광주은행",
        "name": "KJB 동산담보대출_유형자산(원금균등)",
        "account": "1220-028-274781",
        "pay_day": 30,
        "rate": 0.0549,
        "maturity": datetime.date(2026, 8, 30),
        "method": "원금균등분할상환",
        "monthly_principal": 291666000,
        "balance_0531": 874998000,
    },
    {
        "bank": "신협은행",
        "name": "일반한도거래대출금",
        "account": "212-002-616173",
        "pay_day": 24,
        "rate": 0.063,
        "maturity": datetime.date(2027, 1, 24),
        "method": "원리금균등분할",
        "level_payment": 88911825,        # 매월 원리금 균등 납입액
        "balance_0531": 694363893,
    },
]

START = datetime.date(2026, 6, 1)   # 6월분부터 예상


def pay_date(year, month, day):
    """해당 월의 지급일(월말 보정)"""
    import calendar
    last = calendar.monthrange(year, month)[1]
    return datetime.date(year, month, min(day, last))


def month_iter(start_y, start_m, end_date):
    y, m = start_y, start_m
    while True:
        yield y, m
        if (y > end_date.year) or (y == end_date.year and m >= end_date.month):
            break
        m += 1
        if m > 12:
            m = 1
            y += 1


def build_schedule(loan):
    """단일 차입금의 월별 (예정일, 이자, 원금) 스케줄 산출"""
    rows = []
    balance = loan["balance_0531"]
    pd = loan["pay_day"]

    # 직전 지급일(이자 기산일) = 5월 지급일
    prev_date = pay_date(2026, 5, pd)

    for y, m in month_iter(2026, 6, loan["maturity"]):
        cur_date = pay_date(y, m, pd)
        is_last = (y == loan["maturity"].year and m == loan["maturity"].month)

        days = (cur_date - prev_date).days
        interest = round(balance * loan["rate"] * days / DAY_BASIS)

        if loan["method"] == "원리금균등분할":
            if is_last:
                principal = balance
            else:
                principal = loan["level_payment"] - interest
                principal = min(principal, balance)
        else:  # 원금균등 / 자유분할
            if is_last:
                principal = balance
            else:
                principal = min(loan["monthly_principal"], balance)

        rows.append({
            "loan": loan,
            "date": cur_date,
            "interest": interest,
            "principal": principal,
            "balance_after": balance - principal,
        })
        balance -= principal
        prev_date = cur_date
        if balance <= 0:
            break
    return rows


def acct_prefix(account):
    return account.split("-")[0]


# ---------------------------------------------------------------------------
# 전체 스케줄(거래 단위 행) 생성
# ---------------------------------------------------------------------------
records = []
for loan in LOANS:
    for s in build_schedule(loan):
        prefix = acct_prefix(loan["account"])
        # 이자 행
        records.append({
            "date": s["date"],
            "fund": "이자",
            "client": f"{loan['bank'][:2]} {prefix}",
            "bank": loan["bank"],
            "account": loan["account"],
            "summary": f"{loan['account']} 이자",
            "amount": s["interest"],
            "name": loan["name"],
        })
        # 원금 행
        records.append({
            "date": s["date"],
            "fund": "원금",
            "client": f"{loan['bank'][:2]} {prefix}",
            "bank": loan["bank"],
            "account": loan["account"],
            "summary": f"{loan['account']} 원금",
            "amount": s["principal"],
            "name": loan["name"],
        })

# 예정일 → 은행 → 자금과목 순 정렬 (원금 먼저면 이자 다음, 이미지처럼 이자→원금 유지)
fund_order = {"이자": 0, "원금": 1}
records.sort(key=lambda r: (r["date"], r["bank"], fund_order[r["fund"]]))

# ---------------------------------------------------------------------------
# 엑셀 출력 (이미지 컬럼 레이아웃)
# ---------------------------------------------------------------------------
wb = Workbook()
ws = wb.active
ws.title = "이자원금상환스케줄"

headers = ["예정일", "확정일", "자금과목", "거래처명", "사업자(주민)번호",
           "은행", "계좌번호", "예금주", "실제예금", "적요", "금액"]

# 스타일
hdr_fill = PatternFill("solid", fgColor="BDD7EE")
hdr_font = Font(bold=True, size=10)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center")
right = Alignment(horizontal="right", vertical="center")
won_green = PatternFill("solid", fgColor="E2EFDA")  # 원금 행 강조(이미지처럼)

title = "■ NRB 이자/원금 상환 스케줄 (전북은행·신협은행·광주은행) — 2026.06 이후 예상"
ws.cell(1, 1, title).font = Font(bold=True, size=12)
ws.cell(2, 1, "※ 2026-05-31 잔액 기준 / 이자=잔액×이자율×경과일수÷365 / 원금=상환방식에 따른 월 상환액").font = Font(size=9, italic=True, color="808080")

hr = 4
for c, h in enumerate(headers, 1):
    cell = ws.cell(hr, c, h)
    cell.fill = hdr_fill
    cell.font = hdr_font
    cell.alignment = center
    cell.border = border

r = hr + 1
for rec in records:
    is_won = rec["fund"] == "원금"
    vals = [
        rec["date"], rec["date"], rec["fund"], rec["client"], "",
        "", rec["account"], "", "", rec["summary"], rec["amount"],
    ]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(r, c, v)
        cell.border = border
        cell.font = Font(size=10)
        if c in (1, 2):
            cell.number_format = "yyyy-mm-dd"
            cell.alignment = center
        elif c == 3:
            cell.alignment = center
        elif c == 11:
            cell.number_format = "#,##0"
            cell.alignment = right
        else:
            cell.alignment = Alignment(vertical="center")
        if is_won:
            cell.fill = won_green
    r += 1

# 합계 행
last = r
ws.cell(last, 10, "합 계").font = Font(bold=True)
ws.cell(last, 10).alignment = right
total_cell = ws.cell(last, 11, f"=SUM(K{hr+1}:K{last-1})")
total_cell.number_format = "#,##0"
total_cell.font = Font(bold=True)
total_cell.alignment = right
for c in range(1, 12):
    ws.cell(last, c).fill = PatternFill("solid", fgColor="FFF2CC")
    ws.cell(last, c).border = border

# 열 너비
widths = [12, 12, 9, 12, 14, 10, 20, 9, 11, 28, 16]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = "A5"

out = "차입금_이자원금_상환스케줄_전북신협광주.xlsx"
wb.save(out)
print("saved:", out, "rows:", len(records))

# 콘솔 요약
total = sum(x["amount"] for x in records)
print(f"총 거래 행: {len(records)},  총 금액: {total:,}")
by_bank = {}
for x in records:
    by_bank.setdefault(x["bank"], 0)
    by_bank[x["bank"]] += x["amount"]
for b, v in by_bank.items():
    print(f"  {b}: {v:,}")
