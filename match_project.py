"""
입출금 내역의 '품의내역명'과 자금수지의 '적요'를 매칭하여
자금수지의 '프로젝트'를 입출금 내역에 채워넣는 스크립트.

매칭 방식: 자금수지의 '적요' 텍스트가 입출금 내역의 '품의내역명'에 포함되면 매칭
"""

import pandas as pd
from openpyxl import load_workbook
import sys
import os

# ─── 설정 (파일 업로드 후 이 부분 수정) ───────────────────────────────
INPUT_FILE      = "입출금내역_4월.xlsx"   # 입출금 내역 파일명 (NRB 폴더 기준)
CASHFLOW_FILE   = "자금수지계획.xlsx"     # 자금수지 파일명 (NRB 폴더 기준)

INPUT_SHEET     = 0   # 입출금 내역 시트 (0=첫번째, 또는 "Sheet1" 같은 이름)
CASHFLOW_SHEET  = 0   # 자금수지 시트

INPUT_HEADER_ROW    = 0   # 입출금 내역 헤더 행 (0-based: 0=1행, 1=2행 ...)
CASHFLOW_HEADER_ROW = 0   # 자금수지 헤더 행

INPUT_DESC_COL      = "품의내역명"  # 입출금 내역에서 매칭 기준 컬럼
CASHFLOW_DESC_COL   = "적요"        # 자금수지에서 매칭 기준 컬럼
CASHFLOW_PROJECT_COL = "프로젝트"   # 자금수지에서 가져올 컬럼
INPUT_PROJECT_COL   = "프로젝트"    # 입출금 내역에서 채워넣을 컬럼
# ────────────────────────────────────────────────────────────────────────

BASE_DIR = os.path.dirname(os.path.abspath(__file__))


def find_col_letter(ws, header_row, col_name):
    """시트에서 컬럼명으로 열 문자(A, B, ...)를 찾는다."""
    for cell in ws[header_row]:
        if cell.value and str(cell.value).strip() == col_name:
            from openpyxl.utils import get_column_letter
            return get_column_letter(cell.column), cell.column
    return None, None


def match_project(desc, cashflow_df):
    """
    자금수지 '적요'가 입출금 '품의내역명'에 포함되는 행을 찾아 프로젝트명 반환.
    여러 개 매칭 시 첫 번째 값 사용.
    """
    if pd.isna(desc) or not str(desc).strip():
        return None

    desc_str = str(desc).strip()

    # 자금수지 적요가 품의내역명 안에 포함되는지 체크
    mask = cashflow_df[CASHFLOW_DESC_COL].apply(
        lambda x: bool(str(x).strip()) and str(x).strip() in desc_str
        if pd.notna(x) else False
    )
    matched = cashflow_df[mask]

    if matched.empty:
        # 반대 방향: 품의내역명이 적요 안에 포함되는지
        mask2 = cashflow_df[CASHFLOW_DESC_COL].apply(
            lambda x: desc_str in str(x).strip()
            if pd.notna(x) and str(x).strip() else False
        )
        matched = cashflow_df[mask2]

    if not matched.empty:
        return str(matched.iloc[0][CASHFLOW_PROJECT_COL])

    return None


def main():
    input_path    = os.path.join(BASE_DIR, INPUT_FILE)
    cashflow_path = os.path.join(BASE_DIR, CASHFLOW_FILE)

    # 파일 존재 확인
    for path, label in [(input_path, "입출금 내역"), (cashflow_path, "자금수지")]:
        if not os.path.exists(path):
            print(f"[오류] {label} 파일을 찾을 수 없습니다: {path}")
            sys.exit(1)

    # ── 자금수지 읽기 ──────────────────────────────────────────────────
    print(f"자금수지 읽는 중: {CASHFLOW_FILE}")
    cashflow_df = pd.read_excel(
        cashflow_path,
        sheet_name=CASHFLOW_SHEET,
        header=CASHFLOW_HEADER_ROW
    )
    cashflow_df.columns = cashflow_df.columns.astype(str).str.strip()
    print(f"  → {len(cashflow_df)}행")
    print(f"  컬럼: {list(cashflow_df.columns)}")

    for col in [CASHFLOW_DESC_COL, CASHFLOW_PROJECT_COL]:
        if col not in cashflow_df.columns:
            print(f"[오류] 자금수지에서 '{col}' 컬럼을 찾을 수 없습니다.")
            print(f"       사용 가능한 컬럼: {list(cashflow_df.columns)}")
            sys.exit(1)

    # ── 입출금 내역 읽기 ───────────────────────────────────────────────
    print(f"\n입출금 내역 읽는 중: {INPUT_FILE}")
    input_df = pd.read_excel(
        input_path,
        sheet_name=INPUT_SHEET,
        header=INPUT_HEADER_ROW
    )
    input_df.columns = input_df.columns.astype(str).str.strip()
    print(f"  → {len(input_df)}행")
    print(f"  컬럼: {list(input_df.columns)}")

    if INPUT_DESC_COL not in input_df.columns:
        print(f"[오류] 입출금 내역에서 '{INPUT_DESC_COL}' 컬럼을 찾을 수 없습니다.")
        print(f"       사용 가능한 컬럼: {list(input_df.columns)}")
        sys.exit(1)

    # 프로젝트 컬럼이 없으면 추가
    if INPUT_PROJECT_COL not in input_df.columns:
        input_df[INPUT_PROJECT_COL] = None
        print(f"  → '{INPUT_PROJECT_COL}' 컬럼 새로 추가")

    # ── 매칭 ──────────────────────────────────────────────────────────
    print("\n매칭 중...")
    matched_count = 0
    unmatched_rows = []

    for idx, row in input_df.iterrows():
        project = match_project(row[INPUT_DESC_COL], cashflow_df)
        if project:
            input_df.at[idx, INPUT_PROJECT_COL] = project
            matched_count += 1
        else:
            desc = str(row[INPUT_DESC_COL]).strip() if pd.notna(row[INPUT_DESC_COL]) else ""
            if desc:
                unmatched_rows.append((idx + 2, desc))  # 엑셀 행 번호 (헤더 1행 + 1)

    print(f"  → 매칭: {matched_count}건 / 전체: {len(input_df)}건")
    print(f"  → 미매칭: {len(unmatched_rows)}건")

    # ── 결과 저장 ─────────────────────────────────────────────────────
    output_path = input_path.replace(".xlsx", "_프로젝트매칭.xlsx")
    input_df.to_excel(output_path, index=False)
    print(f"\n결과 저장: {os.path.basename(output_path)}")

    # ── 미매칭 목록 출력 ──────────────────────────────────────────────
    if unmatched_rows:
        print("\n[미매칭 항목] 자금수지에서 적요를 찾지 못한 행:")
        print(f"{'행번호':>6}  품의내역명")
        print("-" * 70)
        for row_num, desc in unmatched_rows[:30]:
            print(f"{row_num:>6}  {desc[:65]}")
        if len(unmatched_rows) > 30:
            print(f"         ... 외 {len(unmatched_rows) - 30}건")

    print("\n완료!")


if __name__ == "__main__":
    main()
