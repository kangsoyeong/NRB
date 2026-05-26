import openpyxl

FILE1 = "/root/.claude/uploads/24041d56-69af-4927-919e-4cbd89e59ed0/35a33838-_____260523.xlsx"

wb1 = openpyxl.load_workbook(FILE1, data_only=True)

TARGET = "2026-8778"

print("Sheets in File1:", wb1.sheetnames)
print()

for sheet_name in wb1.sheetnames:
    ws = wb1[sheet_name]
    print(f"--- Sheet: {sheet_name} ---")
    found = False
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        for j, cell in enumerate(row):
            if cell is not None and TARGET in str(cell):
                print(f"  row {i}, col {j}: {repr(cell)}")
                # print surrounding columns
                print(f"  Full row: {row}")
                found = True
    if not found:
        print("  (not found)")
    print()
